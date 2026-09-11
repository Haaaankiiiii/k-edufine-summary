"""K-에듀파인 품목내역 읽기와 공문서 형식 생성. 외부 AI/API를 사용하지 않는다."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import warnings
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLS = 100
MAX_AMOUNT = Decimal("9999999999999999")
FIELDS = ("내용", "규격", "수량", "단위", "예상단가", "예상금액", "수수료")
ALIASES = {
    "내용": {"내용", "품명", "품목명", "물품명", "품목", "적요"},
    "규격": {"규격", "모델명"},
    "수량": {"수량"},
    "단위": {"단위"},
    "예상단가": {"예상단가", "단가"},
    "예상금액": {"예상금액", "금액", "공급금액", "합계금액"},
    "수수료": {"수수료"},
}


class InputError(ValueError):
    pass


def clean(value) -> str:
    if value is None:
        return ""
    # Line breaks in a cell must not become new document instructions or list items.
    return " ".join(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value)).split())


def header_key(value) -> str:
    return re.sub(r"\s|[*＊]|\(원\)|（원）|\[원\]", "", clean(value))


def decimal_value(value, label: str, *, optional=False, whole=False) -> Decimal | None:
    text = clean(value)
    if not text:
        if optional:
            return None
        raise InputError(f"{label}: 값을 입력해 주세요.")
    text = re.sub(r"^[₩￦]", "", text)
    text = re.sub(r"원$", "", text).strip()
    if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text):
        raise InputError(f"{label}: 숫자를 확인해 주세요 ({clean(value)}).")
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise InputError(f"{label}: 올바른 숫자가 아닙니다.") from exc
    if not number.is_finite() or abs(number) > MAX_AMOUNT:
        raise InputError(f"{label}: 지원 범위를 벗어난 숫자입니다.")
    if number < 0:
        raise InputError(f"{label}: 음수는 지원하지 않습니다. 감액 품의는 별도로 작성해 주세요.")
    if whole and number != number.to_integral_value():
        raise InputError(f"{label}: 원 단위 정수로 입력해 주세요. 임의로 반올림하지 않습니다.")
    return number


def number_text(value: Decimal | int) -> str:
    value = Decimal(value)
    return f"{value:,.0f}" if value == value.to_integral_value() else f"{value:,f}".rstrip("0").rstrip(".")


def korean_number(value: int) -> str:
    """Financial writing retains 일 (일십, 일백, 일천, 일만)."""
    if not isinstance(value, int) or value < 0 or value > int(MAX_AMOUNT):
        raise InputError("한글 금액 변환 범위를 벗어났습니다.")
    if value == 0:
        return "영"
    digits, units, groups = "영일이삼사오육칠팔구", ("", "십", "백", "천"), ("", "만", "억", "조")
    parts = []
    for group in groups:
        block = value % 10000
        value //= 10000
        if block:
            segment = "".join(digits[(block // (10 ** i)) % 10] + units[i]
                              for i in range(3, -1, -1) if (block // (10 ** i)) % 10)
            parts.append(segment + group)
    return "".join(reversed(parts))


def budget_text(value: Decimal, suffix=False) -> str:
    if value != value.to_integral_value():
        raise InputError("총액에 원 미만 금액이 있습니다.")
    return f"금{number_text(value)}원(금{korean_number(int(value))}원{'정' if suffix else ''})"


def date_text(value: date, year=True) -> str:
    prefix = f"{value.year}. " if year else ""
    return f"{prefix}{value.month}. {value.day}.({'월화수목금토일'[value.weekday()]})"


def period_text(start: date, end: date) -> str:
    if end < start:
        raise InputError("종료일은 시작일보다 빠를 수 없습니다.")
    if start == end:
        return date_text(start)
    return f"{date_text(start)}~{date_text(end, start.year != end.year)}"


@dataclass
class SheetData:
    name: str
    header_row: int
    rows: list[dict]
    notes: list[str] = field(default_factory=list)
    totals: list[tuple[int, str]] = field(default_factory=list)


def read_workbook(data: bytes) -> list[SheetData]:
    if len(data) > MAX_FILE_BYTES:
        raise InputError("10MB 이하의 .xlsx 파일을 올려 주세요.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise InputError("압축 해제 후 파일 크기가 너무 큽니다. 품목내역 시트만 저장해 주세요.")
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style")
            values = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
            formulas = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except InputError:
        raise
    except Exception as exc:
        raise InputError("엑셀 파일을 읽을 수 없습니다. 암호가 없는 .xlsx 파일인지 확인해 주세요.") from exc

    sheets = []
    try:
        if len(values.sheetnames) > 30:
            raise InputError("시트가 너무 많습니다. 품목내역 시트만 별도 파일로 저장해 주세요.")
        for sheet in values:
            if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLS:
                raise InputError(f"{sheet.title}: 최대 {MAX_ROWS:,}행, {MAX_COLS}열까지 읽을 수 있습니다.")
            raw = list(sheet.iter_rows(values_only=True))
            formula_rows = list(formulas[sheet.title].iter_rows())
            mapping, header_index = {}, -1
            for i, row in enumerate(raw[:50]):
                candidate = {}
                for j, value in enumerate(row):
                    for field_name, aliases in ALIASES.items():
                        if header_key(value) in aliases:
                            if field_name in candidate:
                                raise InputError(f"{sheet.title} {i + 1}행: {field_name} 열이 중복됩니다.")
                            candidate[field_name] = j
                if "내용" in candidate and "예상금액" in candidate:
                    mapping, header_index = candidate, i
                    break
            if header_index < 0:
                continue
            entries, notes, totals = [], [], []
            for i, row in enumerate(raw[header_index + 1:], header_index + 1):
                entry = {key: clean(row[index]) for key, index in mapping.items()}
                entry = {key: entry.get(key, "") for key in FIELDS}
                # Skip only genuinely empty rows; do not silently discard malformed items.
                if not any(entry.values()) and not any(cell.data_type == "f" for cell in formula_rows[i]):
                    continue
                if header_key(entry["내용"]) in ALIASES["내용"] and header_key(entry["예상금액"]) in ALIASES["예상금액"]:
                    continue
                if entry["내용"] in {"합계", "총계", "총합계", "소계"} and not entry["수량"] and not entry["예상단가"]:
                    if entry["내용"] != "소계":
                        totals.append((i + 1, entry["예상금액"]))
                    notes.append(f"엑셀 {i + 1}행의 {entry['내용']}는 품목에서 제외했습니다.")
                    continue
                for name, col in mapping.items():
                    cell = formula_rows[i][col]
                    if cell.data_type == "f" and row[col] is None:
                        entry[name] = "[수식 결과 없음]"
                        notes.append(f"엑셀 {i + 1}행 {name}: 수식 결과가 없어 직접 입력이 필요합니다. 엑셀에서 다시 저장해도 됩니다.")
                entry["원본 행"] = str(i + 1)
                entries.append(entry)
            sheets.append(SheetData(sheet.title, header_index + 1, entries, notes, totals))
    finally:
        values.close()
        formulas.close()
    if not sheets:
        raise InputError("'내용(품명)'과 '예상금액(금액)' 열을 찾지 못했습니다. K-에듀파인 품목내역 .xlsx 파일을 올려 주세요.")
    return sheets


@dataclass(frozen=True)
class Item:
    name: str
    spec: str
    quantity: Decimal | None
    unit: str
    price: Decimal | None
    amount: Decimal
    fee: Decimal
    row: str

    @property
    def matches(self):
        return self.quantity is not None and self.price is not None and self.quantity * self.price == self.amount


@dataclass
class Validation:
    items: list[Item] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    mismatches: list[str] = field(default_factory=list)


def validate_rows(rows: list[dict]) -> Validation:
    result = Validation()
    for i, row in enumerate(rows, 1):
        source = clean(row.get("원본 행")) or str(i)
        label = f"{source}행"
        if not any(clean(row.get(key)) for key in FIELDS):
            continue
        try:
            name = clean(row.get("내용"))
            if not name:
                raise InputError(f"{label}: 품명(내용)이 비어 있습니다.")
            if name == "[수식 결과 없음]":
                raise InputError(f"{label}: 품명 수식의 결과를 입력해 주세요.")
            qty = decimal_value(row.get("수량"), f"{label} 수량", optional=True)
            price = decimal_value(row.get("예상단가"), f"{label} 예상단가", optional=True)
            amount = decimal_value(row.get("예상금액"), f"{label} 예상금액", optional=True, whole=True)
            fee = decimal_value(row.get("수수료"), f"{label} 수수료", optional=True, whole=True) or Decimal(0)
            if qty is not None and qty == 0:
                raise InputError(f"{label}: 수량은 0보다 커야 합니다.")
            if amount is None:
                if qty is None or price is None:
                    raise InputError(f"{label}: 예상금액 또는 수량·예상단가를 입력해 주세요.")
                amount = qty * price
                if amount != amount.to_integral_value():
                    raise InputError(f"{label}: 계산값에 원 미만 금액이 있습니다. 예상금액을 직접 입력해 주세요.")
                result.notes.append(f"{label}: 빈 예상금액을 단가 × 수량으로 계산했습니다 ({number_text(amount)}원).")
            if amount > MAX_AMOUNT:
                raise InputError(f"{label}: 예상금액이 지원 범위를 벗어났습니다.")
            item = Item(name, clean(row.get("규격")), qty, clean(row.get("단위")), price, amount, fee, source)
            if qty is not None and price is not None and not item.matches:
                result.mismatches.append(f"{label} {name}: 단가 × 수량은 {number_text(qty * price)}원, 예상금액은 {number_text(amount)}원입니다.")
            if qty is None or price is None:
                result.notes.append(f"{label}: 수량 또는 단가가 없어 예상금액만 표기합니다.")
            if qty is not None and not item.unit:
                result.notes.append(f"{label}: 단위가 비어 있어 숫자 '{number_text(qty)}'만 표기합니다. 위 표에서 개·권·세트 등을 입력할 수 있습니다.")
            result.items.append(item)
        except InputError as exc:
            result.errors.append(str(exc))
    if not result.items and not result.errors:
        result.errors.append("작성할 품목이 없습니다.")
    return result


def item_name(item: Item, show_spec=True) -> str:
    return item.name + (f"(규격: {item.spec})" if show_spec and item.spec else "")


def detail_line(item: Item, *, show_spec=True, fee_mode="none", allow_mismatch=False) -> str:
    prefix = item_name(item, show_spec) + ": "
    if item.quantity is not None and item.price is not None and not item.matches and not allow_mismatch:
        raise InputError("단가 × 수량과 예상금액의 차이를 확인해 주세요.")
    if item.matches:
        qty = f"{number_text(item.quantity)}{item.unit}"
        body = f"{number_text(item.price)}원 X {qty} = {number_text(item.amount)}원"
    else:
        body = f"{number_text(item.amount)}원(예상금액 기준)"
    if item.fee and fee_mode == "add":
        body += f", 수수료 {number_text(item.fee)}원, 합계 {number_text(item.amount + item.fee)}원"
    elif item.fee and fee_mode == "included":
        body += f"(수수료 {number_text(item.fee)}원 포함)"
    return prefix + body


def item_lines(lines: list[str]) -> list[str]:
    # The source explicitly continues 하 with 거. For very long lists a third-level
    # numbered list avoids unsupported or duplicated Korean item symbols.
    consonants = [0, 2, 3, 5, 6, 7, 9, 11, 12, 14, 15, 16, 17, 18]
    vowels = [0, 4, 8, 13, 18, 20]
    labels = [chr(0xAC00 + c * 588 + v * 28) for v in vowels for c in consonants]
    if len(lines) <= len(labels):
        return [f"  {labels[i]}. {line}" for i, line in enumerate(lines)]
    return ["  가. 품목별 내역"] + [f"    {i}) {line}" for i, line in enumerate(lines, 1)]


def total_amount(items: list[Item], fee_mode="none") -> Decimal:
    if fee_mode not in {"none", "add", "included", "exclude"}:
        raise InputError("수수료 처리 방식을 선택해 주세요.")
    if any(item.fee for item in items) and fee_mode == "none":
        raise InputError("수수료 처리 방식을 선택해 주세요.")
    if fee_mode == "included" and any(item.fee > item.amount for item in items):
        raise InputError("포함 수수료가 예상금액보다 큰 품목이 있습니다.")
    total = sum((item.amount + (item.fee if fee_mode == "add" else 0) for item in items), Decimal(0))
    if total > MAX_AMOUNT:
        raise InputError("총액이 한글 금액 변환 범위를 벗어났습니다.")
    return total


@dataclass
class DocumentOptions:
    purpose: str
    subject: str = "물품"
    action: str = "구입"
    related: str = ""
    period: str = ""
    place: str = ""
    budget_account: str = ""
    attachments: str = ""
    show_spec: bool = True
    all_names: bool = False
    suffix: bool = False
    fee_mode: str = "none"
    allow_mismatch: bool = False


def generate_document(items: list[Item], options: DocumentOptions) -> tuple[str, str, Decimal]:
    purpose, subject = clean(options.purpose), clean(options.subject)
    if not purpose or not subject:
        raise InputError("목적과 구입·지급 대상을 입력해 주세요.")
    if not items:
        raise InputError("작성할 품목이 없습니다.")
    total = total_amount(items, options.fee_mode)
    title = f"{subject} {options.action}"
    last = subject[-1]
    if "가" <= last <= "힣":
        particle = "을" if (ord(last) - 0xAC00) % 28 else "를"
        opening = f"{subject}{particle} 아래와 같이 {options.action}하고자 합니다."
    else:
        opening = f"다음과 같이 {subject} {options.action}을 하고자 합니다."
    lines = [opening]
    sections = []
    if clean(options.related):
        sections.append(("관련", clean(options.related)))
    sections.append(("목적", purpose))
    if options.all_names:
        names = "; ".join(item_name(item, options.show_spec) for item in items)
    else:
        first = item_name(items[0], options.show_spec)
        names = first if len(items) == 1 else f"{first} 외 {len(items) - 1}건(총 {len(items)}건, 산출 근거 참조)"
    sections.append(("품명", names))
    if options.period:
        sections.append(("기간", options.period))
    if clean(options.place):
        sections.append(("사용 장소", clean(options.place)))
    sections.append(("소요 예산", budget_text(total, options.suffix)))
    if clean(options.budget_account):
        sections.append(("예산 과목", clean(options.budget_account)))
    sections.append(("산출 근거", ""))
    lines.extend(f"{i}. {name}: {body}".rstrip() for i, (name, body) in enumerate(sections, 1))
    details = [detail_line(item, show_spec=options.show_spec, fee_mode=options.fee_mode,
                           allow_mismatch=options.allow_mismatch) for item in items]
    lines.extend(item_lines(details))
    attachments = [clean(line).rstrip(".") for line in options.attachments.splitlines() if clean(line)]
    if attachments:
        lines.append("")
        if len(attachments) == 1:
            lines.append(f"붙임  {attachments[0]}.  끝.")
        else:
            lines.append(f"붙임  1. {attachments[0]}.")
            lines.extend(f"      {i}. {name}." for i, name in enumerate(attachments[1:], 2))
            lines[-1] += "  끝."
    else:
        lines[-1] += "  끝."
    return title, "\n".join(lines), total


def demo_rows() -> list[dict]:
    return [
        dict(zip(("원본 행", *FIELDS), ("4", "A4 복사용지", "80g, 2,500매", "5", "상자", "24000", "120000", "0"))),
        dict(zip(("원본 행", *FIELDS), ("5", "보드마카", "검정, 12개입", "3", "세트", "12000", "36000", "0"))),
        dict(zip(("원본 행", *FIELDS), ("6", "클리어파일", "A4, 40매", "10", "권", "3000", "30000", "0"))),
    ]
